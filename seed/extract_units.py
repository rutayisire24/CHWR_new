#!/usr/bin/env python3
"""Extract the national hierarchy to TSV for loading by seed/load_hierarchy.sql.

Sources:
  data/Village-Admin Units 06-08-2026.xlsm   districts downward (authoritative)
  <ODK workbook>                             the 15 regions and their district map

Usage: python3 seed/extract_units.py <odk_workbook.xlsx> [out_dir]   (default seed/out)
"""
import openpyxl, csv, re, sys, os

norm = lambda s: re.sub(r'[^a-z0-9]', '', (s or '').lower())
HERE = os.path.dirname(os.path.abspath(__file__))
UNITS = os.path.join(HERE, "..", "data", "Village-Admin Units 06-08-2026.xlsm")


def main(odk_path, out):
    os.makedirs(out, exist_ok=True)

    ws = openpyxl.load_workbook(UNITS, read_only=True, data_only=True).active
    it = ws.iter_rows(values_only=True); next(it)
    n = 0
    with open(os.path.join(out, "units.tsv"), "w", newline="") as f:
        w = csv.writer(f, delimiter="\t", quoting=csv.QUOTE_NONE, escapechar="\\")
        for r in it:
            if not r or r[0] in (None, ""):
                continue
            w.writerow([str(x).strip() if x is not None else "" for x in r[:10]])
            n += 1

    wb = openpyxl.load_workbook(odk_path, read_only=True, data_only=True)["choices"]
    it2 = wb.iter_rows(values_only=True)
    hdr = [str(x).strip() if x else "" for x in next(it2)]
    J = {k: i for i, k in enumerate(hdr)}
    regions, dmap = {}, {}
    for r in it2:
        if not r or not r[0]:
            continue
        ln = str(r[0]).strip()
        if ln == "region":
            regions[str(r[J['name']]).strip()] = str(r[J['label']]).strip()
        elif ln == "district":
            dmap[norm(str(r[J['label']]))] = str(r[J['regionfilter']]).strip()

    with open(os.path.join(out, "regions.tsv"), "w") as f:
        for k, v in regions.items():
            f.write(f"{k}\t{v}\n")
    with open(os.path.join(out, "dregion.tsv"), "w") as f:
        for k, v in dmap.items():
            f.write(f"{k}\t{v}\n")

    print(f"units={n} regions={len(regions)} district_region_map={len(dmap)}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "out"))
